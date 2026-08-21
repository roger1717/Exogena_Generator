"""
Procesador de archivos para lectura de CSV y Excel

Este módulo proporciona funciones para leer archivos CSV y Excel
de manera consistente, manejando diferentes formatos y codificaciones.


"""

import csv
import json
from pathlib import Path
from typing import List, Dict, Any, Optional, Union
from datetime import datetime
import openpyxl
from openpyxl import load_workbook
from loguru import logger


class FileProcessor:
    """
    Clase para manejar diferentes formatos de archivos
    """
    
    @staticmethod
    def leer_csv(
        file_path: Union[str, Path],
        delimiter: str = ',',
        encoding: str = 'utf-8-sig',
        has_header: bool = True
    ) -> List[Dict[str, Any]]:
        """
        Leer un archivo CSV y retornar como lista de diccionarios
        
        Args:
            file_path: Ruta del archivo CSV
            delimiter: Delimitador de campos
            encoding: Codificación del archivo
            has_header: Si tiene encabezados
        
        Returns:
            List[Dict[str, Any]]: Datos del CSV
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"Archivo no encontrado: {file_path}")
        
        logger.info(f"Leyendo CSV: {file_path}")
        datos = []
        
        with open(file_path, 'r', encoding=encoding) as file:
            if has_header:
                reader = csv.DictReader(file, delimiter=delimiter)
                datos = list(reader)
            else:
                reader = csv.reader(file, delimiter=delimiter)
                for row in reader:
                    datos.append({
                        f"col_{i}": value for i, value in enumerate(row)
                    })
        
        logger.info(f"CSV leído: {len(datos)} registros")
        return datos
    
    @staticmethod
    def escribir_csv(
        datos: List[Dict[str, Any]],
        file_path: Union[str, Path],
        delimiter: str = ',',
        encoding: str = 'utf-8-sig'
    ) -> None:
        """
        Escribir datos a un archivo CSV
        
        Args:
            datos: Lista de diccionarios con los datos
            file_path: Ruta del archivo de salida
            delimiter: Delimitador de campos
            encoding: Codificación del archivo
        """
        file_path = Path(file_path)
        file_path.parent.mkdir(parents=True, exist_ok=True)
        
        if not datos:
            logger.warning("No hay datos para escribir")
            return
        
        # Obtener encabezados de la primera fila
        headers = list(datos[0].keys())
        
        logger.info(f"Escribiendo CSV: {file_path} ({len(datos)} registros)")
        
        with open(file_path, 'w', encoding=encoding, newline='') as file:
            writer = csv.DictWriter(file, fieldnames=headers, delimiter=delimiter)
            writer.writeheader()
            writer.writerows(datos)
    
    @staticmethod
    def leer_excel(
        file_path: Union[str, Path],
        sheet_name: Optional[str] = None,
        has_header: bool = True
    ) -> List[Dict[str, Any]]:
        """
        Leer un archivo Excel y retornar como lista de diccionarios
        
        Args:
            file_path: Ruta del archivo Excel
            sheet_name: Nombre de la hoja (None para la primera)
            has_header: Si tiene encabezados
        
        Returns:
            List[Dict[str, Any]]: Datos del Excel
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"Archivo no encontrado: {file_path}")
        
        logger.info(f"Leyendo Excel: {file_path}")
        
        # Cargar el workbook
        wb = load_workbook(file_path, data_only=True)
        
        # Seleccionar hoja
        if sheet_name:
            ws = wb[sheet_name]
        else:
            ws = wb.active
        
        # Leer datos
        datos = []
        rows = list(ws.iter_rows(values_only=True))
        
        if not rows:
            wb.close()
            return datos
        
        # Obtener encabezados
        if has_header:
            headers = []
            for cell in rows[0]:
                if cell is None:
                    headers.append(f"col_{len(headers)}")
                else:
                    headers.append(str(cell).strip())
            data_rows = rows[1:]
        else:
            headers = [f"col_{i}" for i in range(len(rows[0]))]
            data_rows = rows
        
        # Convertir a diccionarios
        for row in data_rows:
            # Omitir filas completamente vacías
            if all(cell is None or str(cell).strip() == '' for cell in row):
                continue
            
            row_dict = {}
            for i, cell in enumerate(row):
                if i < len(headers):
                    # Convertir tipos para consistencia
                    if isinstance(cell, datetime):
                        row_dict[headers[i]] = cell.strftime('%Y-%m-%d')
                    elif isinstance(cell, (int, float)):
                        row_dict[headers[i]] = float(cell)
                    elif cell is None:
                        row_dict[headers[i]] = None
                    else:
                        row_dict[headers[i]] = str(cell).strip()
            
            # Solo agregar si hay al menos un valor no vacío
            if any(v is not None and v != '' for v in row_dict.values()):
                datos.append(row_dict)
        
        wb.close()
        logger.info(f"Excel leído: {len(datos)} registros")
        return datos
    
    @staticmethod
    def escribir_excel(
        datos: List[Dict[str, Any]],
        file_path: Union[str, Path],
        sheet_name: str = "Datos",
        headers: Optional[List[str]] = None
    ) -> None:
        """
        Escribir datos a un archivo Excel
        
        Args:
            datos: Lista de diccionarios con los datos
            file_path: Ruta del archivo de salida
            sheet_name: Nombre de la hoja
            headers: Lista de encabezados (si None, usa las keys del primer diccionario)
        """
        file_path = Path(file_path)
        file_path.parent.mkdir(parents=True, exist_ok=True)
        
        if not datos:
            logger.warning("No hay datos para escribir")
            return
        
        # Determinar encabezados
        if headers is None:
            headers = list(datos[0].keys())
        
        logger.info(f"Escribiendo Excel: {file_path} ({len(datos)} registros)")
        
        # Crear workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        
        # Escribir encabezados
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=header)
        
        # Escribir datos
        for row_idx, row_data in enumerate(datos, 2):
            for col_idx, header in enumerate(headers, 1):
                value = row_data.get(header, None)
                # Convertir fechas a string si es necesario
                if isinstance(value, datetime):
                    value = value.strftime('%Y-%m-%d')
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        # Ajustar ancho de columnas
        for col_idx, header in enumerate(headers, 1):
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            max_length = len(str(header))
            for row_idx in range(2, min(len(datos) + 2, 100)):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_length + 2, 50)
        
        # Guardar
        wb.save(file_path)
        wb.close()
    
    @staticmethod
    def leer_json(
        file_path: Union[str, Path],
        encoding: str = 'utf-8'
    ) -> Union[List[Dict[str, Any]], Dict[str, Any]]:
        """
        Leer un archivo JSON
        
        Args:
            file_path: Ruta del archivo JSON
            encoding: Codificación del archivo
        
        Returns:
            Datos del JSON (dict o list)
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"Archivo no encontrado: {file_path}")
        
        with open(file_path, 'r', encoding=encoding) as file:
            datos = json.load(file)
        
        return datos
    
    @staticmethod
    def escribir_json(
        datos: Any,
        file_path: Union[str, Path],
        indent: int = 2,
        encoding: str = 'utf-8'
    ) -> None:
        """
        Escribir datos a un archivo JSON
        
        Args:
            datos: Datos a escribir
            file_path: Ruta del archivo de salida
            indent: Indentación para formateo
            encoding: Codificación del archivo
        """
        file_path = Path(file_path)
        file_path.parent.mkdir(parents=True, exist_ok=True)
        
        with open(file_path, 'w', encoding=encoding) as file:
            json.dump(datos, file, indent=indent, ensure_ascii=False)
    
    @staticmethod
    def detectar_formato(file_path: Union[str, Path]) -> str:
        """
        Detectar el formato de un archivo por su extensión
        
        Args:
            file_path: Ruta del archivo
        
        Returns:
            str: 'csv', 'excel', 'json' o 'desconocido'
        """
        file_path = Path(file_path)
        ext = file_path.suffix.lower()
        
        if ext in ['.csv']:
            return 'csv'
        elif ext in ['.xlsx', '.xls']:
            return 'excel'
        elif ext in ['.json']:
            return 'json'
        else:
            return 'desconocido'
    
    @staticmethod
    def leer_archivo(
        file_path: Union[str, Path],
        **kwargs
    ) -> List[Dict[str, Any]]:
        """
        Leer un archivo detectando automáticamente su formato
        
        Args:
            file_path: Ruta del archivo
            **kwargs: Argumentos adicionales para cada lector
        
        Returns:
            List[Dict[str, Any]]: Datos del archivo
        """
        file_path = Path(file_path)
        formato = FileProcessor.detectar_formato(file_path)
        
        if formato == 'csv':
            return FileProcessor.leer_csv(file_path, **kwargs)
        elif formato == 'excel':
            return FileProcessor.leer_excel(file_path, **kwargs)
        elif formato == 'json':
            resultado = FileProcessor.leer_json(file_path, **kwargs)
            if isinstance(resultado, list):
                return resultado
            else:
                return [resultado]
        else:
            raise ValueError(f"Formato de archivo no soportado: {file_path.suffix}")