#!/usr/bin/env python

"""
Generador de archivos Excel para Retenciones en la Fuente

Autor: Roger Hoyos
Fecha: 2026-08-22
Versión: 2.0.0
"""

import logging
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

logger = logging.getLogger(__name__)

# Estilos globales para reutilizar en el reporte
BORDE_FINO = Side(border_style="thin", color="D9D9D9")
BORDE = Border(left=BORDE_FINO, right=BORDE_FINO, top=BORDE_FINO, bottom=BORDE_FINO)


def _set_header_at(ws, row_idx: int, headers: list):
    """Función auxiliar para escribir y formatear encabezados de tablas."""
    fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    font = Font(bold=True, color="000000")
    alignment = Alignment(horizontal="center", vertical="center")

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=header)
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment


class ExcelGenerator:

    TIPO_A_NOMBRE = {
        "renta": "ReteFuente",
        "iva": "ReteIVA",
        "ica": "ReteICA",
    }

    def __init__(self, output_dir: Path = Path("outputs/excel")):
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generar_reporte_retenciones(
        self,
        retenciones: List[Dict],
        periodo: Optional[str] = None,
        tipo: Optional[str] = None,
    ) -> Path:
        """Generar reporte de retenciones en Excel"""

        if not retenciones:
            raise ValueError("No hay retenciones para generar")

        if tipo is None and retenciones:
            tipo = retenciones[0].get("tipo_retencion", "renta")

        nombre_tipo = self.TIPO_A_NOMBRE.get(tipo, "Retenciones")

        if periodo:
            filename = f"1003-{nombre_tipo}-{periodo}.xlsx"
        else:
            filename = f"1003-{nombre_tipo}-{datetime.now().strftime('%Y%m%d')}.xlsx"

        file_path = self.output_dir / filename
        wb = Workbook()

        # === HOJA 1: Detalle ===
        ws_detalle = wb.active
        ws_detalle.title = "Detalle"

        df = pd.DataFrame(retenciones)

        columnas = [
            ("fecha", "Fecha"),
            ("comprobante", "Comprobante"),
            ("nit_tercero", "NIT"),
            ("nombre_tercero", "Nombre Tercero"),
            ("concepto_contable", "Concepto Contable"),
            ("concepto_dian", "Código DIAN"),
            ("base_gravable", "Base Gravable"),
            ("tarifa", "Tarifa"),
            ("valor_retenido", "Valor Retenido"),
            ("cuenta_pasivo", "Cuenta Pasivo"),
            ("periodo", "Período"),
            ("tipo_retencion", "Tipo"),
        ]

        columnas_existentes = [(orig, dest) for orig, dest in columnas if orig in df.columns]
        df_detalle = df[[orig for orig, _ in columnas_existentes]].copy()
        df_detalle.columns = [dest for _, dest in columnas_existentes]

        for col_idx, col_name in enumerate(df_detalle.columns, 1):
            cell = ws_detalle.cell(row=1, column=col_idx, value=col_name)
            cell.font = Font(bold=True, color="000000")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        for r_idx, row in enumerate(df_detalle.itertuples(index=False), 2):
            for c_idx, value in enumerate(row, 1):
                cell = ws_detalle.cell(row=r_idx, column=c_idx, value=value)
                if c_idx in [7, 9]:
                    cell.number_format = "#,##0.00"
                if c_idx == 8:
                    cell.number_format = "0.00%"

        for col in ws_detalle.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws_detalle.column_dimensions[column].width = adjusted_width

        # === HOJA 2: Resumen por Concepto ===
        ws_resumen = wb.create_sheet("Resumen por Concepto")

        if "concepto_dian" in df.columns:
            from app.core.constants import CONCEPTOS_RETENCION_DIAN

            df_resumen = df.groupby("concepto_dian").agg(
                {"valor_retenido": ["count", "sum"], "base_gravable": "sum"}
            ).round(2)

            df_resumen.columns = ["Cantidad", "Total Retenido", "Total Base"]
            df_resumen = df_resumen.reset_index()
            df_resumen["Concepto"] = df_resumen["concepto_dian"].map(
                lambda x: CONCEPTOS_RETENCION_DIAN.get(str(x), x)
            )
            df_resumen = df_resumen[["concepto_dian", "Concepto", "Cantidad", "Total Base", "Total Retenido"]]

            for col_idx, col_name in enumerate(df_resumen.columns, 1):
                cell = ws_resumen.cell(row=1, column=col_idx, value=col_name)
                cell.font = Font(bold=True, color="000000")
                cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")

            for r_idx, row in enumerate(df_resumen.itertuples(index=False), 2):
                for c_idx, value in enumerate(row, 1):
                    cell = ws_resumen.cell(row=r_idx, column=c_idx, value=value)
                    if c_idx in [4, 5]:
                        cell.number_format = "#,##0.00"

            for col in ws_resumen.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                adjusted_width = min(max_length + 2, 30)
                ws_resumen.column_dimensions[column].width = adjusted_width

        # === HOJA 3: Resumen por Período ===
        ws_periodo = wb.create_sheet("Resumen por Período")

        if "periodo" in df.columns:
            df_periodo = df.groupby("periodo").agg(
                {"valor_retenido": ["count", "sum"]}
            ).round(2)
            df_periodo.columns = ["Cantidad", "Total Retenido"]
            df_periodo = df_periodo.reset_index()

            for col_idx, col_name in enumerate(df_periodo.columns, 1):
                cell = ws_periodo.cell(row=1, column=col_idx, value=col_name)
                cell.font = Font(bold=True, color="000000")
                cell.fill = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")

            for r_idx, row in enumerate(df_periodo.itertuples(index=False), 2):
                for c_idx, value in enumerate(row, 1):
                    cell = ws_periodo.cell(row=r_idx, column=c_idx, value=value)
                    if c_idx == 3:
                        cell.number_format = "#,##0.00"

            for col in ws_periodo.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                adjusted_width = min(max_length + 2, 30)
                ws_periodo.column_dimensions[column].width = adjusted_width

        # === HOJA 4: Resumen por Tipo ===
        ws_tipo = wb.create_sheet("Resumen por Tipo")

        if "tipo_retencion" in df.columns:
            df_tipo = df.groupby("tipo_retencion").agg(
                {"valor_retenido": ["count", "sum"]}
            ).round(2)
            df_tipo.columns = ["Cantidad", "Total Retenido"]
            df_tipo = df_tipo.reset_index()
            df_tipo["Tipo"] = df_tipo["tipo_retencion"].map(
                lambda x: self.TIPO_A_NOMBRE.get(x, x)
            )
            df_tipo = df_tipo[["Tipo", "Cantidad", "Total Retenido"]]

            for col_idx, col_name in enumerate(df_tipo.columns, 1):
                cell = ws_tipo.cell(row=1, column=col_idx, value=col_name)
                cell.font = Font(bold=True, color="000000")
                cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")

            for r_idx, row in enumerate(df_tipo.itertuples(index=False), 2):
                for c_idx, value in enumerate(row, 1):
                    cell = ws_tipo.cell(row=r_idx, column=c_idx, value=value)
                    if c_idx == 3:
                        cell.number_format = "#,##0.00"

            for col in ws_tipo.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                adjusted_width = min(max_length + 2, 30)
                ws_tipo.column_dimensions[column].width = adjusted_width

        wb.save(file_path)
        logger.info(f"📊 Excel generado: {file_path}")

        return file_path

    def generar_reporte_por_tipo(
        self,
        retenciones: List[Dict],
        periodo: Optional[str] = None
    ) -> Dict[str, Path]:
        """Generar reportes separados por tipo de retención"""
        if not retenciones:
            return {}

        retenciones_por_tipo = {}
        for r in retenciones:
            tipo = r.get("tipo_retencion", "renta")
            if tipo not in retenciones_por_tipo:
                retenciones_por_tipo[tipo] = []
            retenciones_por_tipo[tipo].append(r)

        archivos_generados = {}

        for tipo, retenciones_tipo in retenciones_por_tipo.items():
            archivo = self.generar_reporte_retenciones(
                retenciones=retenciones_tipo,
                periodo=periodo,
                tipo=tipo
            )
            archivos_generados[tipo] = archivo

        return archivos_generados

    def generar_reporte_unificado(
        self,
        retenciones: List[Dict],
        periodo: Optional[str] = None
    ) -> Path:
        """
        Genera un único archivo Excel con TODAS las retenciones consolidadas.
        Útil para la declaración de renta (certificado de retenciones).
        """
        if not retenciones:
            raise ValueError("No hay retenciones para generar el reporte unificado")

        # Nombre del archivo (formato 1003 - Retenciones Certificadas)
        if periodo:
            filename = f"1003-CertificadasRetenciones-{periodo}.xlsx"
        else:
            filename = f"1003-CertificadasRetenciones-{datetime.now().strftime('%Y%m%d')}.xlsx"

        file_path = self.output_dir / filename
        wb = Workbook()

        # --- HOJA 1: RESUMEN GENERAL (Totales por Tipo) ---
        ws_resumen = wb.active
        ws_resumen.title = "Resumen General"

        df = pd.DataFrame(retenciones)

        # Calcular totales
        total_general = df["valor_retenido"].sum()

        # Encabezados del resumen
        ws_resumen["A1"] = "CERTIFICADO DE RETENCIONES EN LA FUENTE"
        ws_resumen["A1"].font = Font(bold=True, size=14, color="1F4E78")
        ws_resumen["A2"] = f"Período: {periodo if periodo else 'General'}"
        ws_resumen["A2"].font = Font(italic=True, color="595959")
        ws_resumen["A3"] = f"Total registros: {len(df)}"

        # Tabla de resumen por tipo
        headers_resumen = ["Tipo de Retención", "Cantidad", "Total Retenido"]
        _set_header_at(ws_resumen, 5, headers_resumen)

        tipos = [
            ("ReteFuente (Renta)", "renta"),
            ("ReteIVA", "iva"),
            ("ReteICA", "ica"),
            ("TOTAL GENERAL", None),
        ]

        row = 6
        for nombre, tipo in tipos:
            if tipo is None:
                ws_resumen.cell(row=row, column=1, value=nombre).font = Font(bold=True)
                ws_resumen.cell(row=row, column=2, value=len(df)).font = Font(bold=True)
                cell = ws_resumen.cell(row=row, column=3, value=total_general)
                cell.number_format = "#,##0.00"
                cell.font = Font(bold=True)
                for col_idx in range(1, 4):
                    c = ws_resumen.cell(row=row, column=col_idx)
                    c.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                    c.font = Font(bold=True, color="000000")
            else:
                count = len(df[df["tipo_retencion"] == tipo])
                total = df[df["tipo_retencion"] == tipo]["valor_retenido"].sum()
                ws_resumen.cell(row=row, column=1, value=nombre)
                ws_resumen.cell(row=row, column=2, value=count)
                cell = ws_resumen.cell(row=row, column=3, value=total)
                cell.number_format = "#,##0.00"
            row += 1

        for col in ["A", "B", "C"]:
            ws_resumen.column_dimensions[col].width = 25

        # --- HOJA 2: DETALLE UNIFICADO ---
        ws_detalle = wb.create_sheet("Detalle Unificado")

        df["tipo_nombre"] = df["tipo_retencion"].map(self.TIPO_A_NOMBRE)

        columnas = [
            ("fecha", "Fecha"),
            ("comprobante", "Comprobante"),
            ("nit_tercero", "NIT"),
            ("nombre_tercero", "Nombre Tercero"),
            ("concepto_contable", "Concepto Contable"),
            ("concepto_dian", "Código DIAN"),
            ("base_gravable", "Base Gravable"),
            ("tarifa", "Tarifa"),
            ("valor_retenido", "Valor Retenido"),
            ("cuenta_pasivo", "Cuenta Pasivo"),
            ("periodo", "Período"),
            ("tipo_nombre", "Tipo Retención"),
        ]

        columnas_existentes = [(orig, dest) for orig, dest in columnas if orig in df.columns]
        df_detalle = df[[orig for orig, _ in columnas_existentes]].copy()
        df_detalle.columns = [dest for _, dest in columnas_existentes]

        _set_header_at(ws_detalle, 1, df_detalle.columns.tolist())

        for r_idx, row_data in enumerate(df_detalle.itertuples(index=False), 2):
            for c_idx, value in enumerate(row_data, 1):
                cell = ws_detalle.cell(row=r_idx, column=c_idx, value=value)
                cell.border = BORDE
                if c_idx in [7, 9]:
                    cell.number_format = "#,##0.00"
                if c_idx == 8:
                    cell.number_format = "0.00%"

        for col in ws_detalle.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws_detalle.column_dimensions[column].width = adjusted_width

        # --- HOJA 3: RESUMEN PERÍODO Y TIPO ---
        ws_periodo = wb.create_sheet("Resumen Período-Tipo")

        if "periodo" in df.columns and "tipo_retencion" in df.columns:
            df_periodo = df.groupby(["periodo", "tipo_retencion"]).agg(
                {"valor_retenido": ["count", "sum"]}
            ).round(2)
            df_periodo.columns = ["Cantidad", "Total Retenido"]
            df_periodo = df_periodo.reset_index()
            df_periodo["Tipo"] = df_periodo["tipo_retencion"].map(self.TIPO_A_NOMBRE)
            df_periodo = df_periodo[["periodo", "Tipo", "Cantidad", "Total Retenido"]]

            _set_header_at(ws_periodo, 1, df_periodo.columns.tolist())

            for r_idx, row_data in enumerate(df_periodo.itertuples(index=False), 2):
                for c_idx, value in enumerate(row_data, 1):
                    cell = ws_periodo.cell(row=r_idx, column=c_idx, value=value)
                    if c_idx == 4:
                        cell.number_format = "#,##0.00"

            for col in ws_periodo.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                adjusted_width = min(max_length + 2, 30)
                ws_periodo.column_dimensions[column].width = adjusted_width

        # --- HOJA 4: RESUMEN POR CONCEPTO DIAN ---
        ws_concepto = wb.create_sheet("Resumen por Concepto")

        if "concepto_dian" in df.columns:
            from app.core.constants import CONCEPTOS_RETENCION_DIAN

            df_concepto = df.groupby("concepto_dian").agg(
                {"valor_retenido": ["count", "sum"], "base_gravable": "sum"}
            ).round(2)
            df_concepto.columns = ["Cantidad", "Total Retenido", "Total Base"]
            df_concepto = df_concepto.reset_index()
            df_concepto["Concepto"] = df_concepto["concepto_dian"].map(
                lambda x: CONCEPTOS_RETENCION_DIAN.get(str(x), x)
            )
            df_concepto = df_concepto[["concepto_dian", "Concepto", "Cantidad", "Total Base", "Total Retenido"]]

            _set_header_at(ws_concepto, 1, df_concepto.columns.tolist())

            for r_idx, row_data in enumerate(df_concepto.itertuples(index=False), 2):
                for c_idx, value in enumerate(row_data, 1):
                    cell = ws_concepto.cell(row=r_idx, column=c_idx, value=value)
                    if c_idx in [4, 5]:
                        cell.number_format = "#,##0.00"

            for col in ws_concepto.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                adjusted_width = min(max_length + 2, 30)
                ws_concepto.column_dimensions[column].width = adjusted_width

        wb.save(file_path)
        logger.info(f"📊 Certificado de retenciones generado: {file_path}")
        return file_path