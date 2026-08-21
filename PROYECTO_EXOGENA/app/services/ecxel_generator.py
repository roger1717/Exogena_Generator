# app/services/excel_generator.py

"""
Generador de archivos Excel para Retenciones y Exógena

Autor: [Tu nombre]
Fecha: 2026-08-21
Versión: 1.0.0
"""

import pandas as pd
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from typing import List, Dict, Optional


class ExcelGenerator:
    """
    Generador de archivos Excel para retenciones y exógena
    """
    
    def __init__(self, output_dir: Path = Path("outputs/excel")):
        """
        Inicializar el generador
        
        Args:
            output_dir: Directorio de salida para los archivos Excel
        """
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)
    
    def generar_reporte_retenciones(
        self,
        retenciones: List[Dict],
        periodo: Optional[str] = None
    ) -> Path:
        """
        Generar reporte de retenciones en Excel
        
        Args:
            retenciones: Lista de retenciones procesadas
            periodo: Período (opcional, para el nombre del archivo)
        
        Returns:
            Path del archivo generado
        """
        if not retenciones:
            raise ValueError("No hay retenciones para generar")
        
        # Crear nombre del archivo
        fecha = datetime.now().strftime("%Y%m%d")
        periodo_str = f"_{periodo}" if periodo else ""
        filename = f"{fecha}_ReteFuente{periodo_str}.xlsx"
        file_path = self.output_dir / filename
        
        # Crear workbook
        wb = Workbook()
        
        # --- Hoja 1: Detalle de Retenciones ---
        ws_detalle = wb.active
        ws_detalle.title = "Retenciones"
        
        # Convertir a DataFrame
        df = pd.DataFrame(retenciones)
        
        # Seleccionar y renombrar columnas
        columnas = [
            ('fecha', 'Fecha'),
            ('comprobante', 'Comprobante'),
            ('nit_tercero', 'NIT'),
            ('nombre_tercero', 'Nombre Tercero'),
            ('concepto_contable', 'Concepto Contable'),
            ('concepto_dian', 'Concepto DIAN'),
            ('base_gravable', 'Base Gravable'),
            ('tarifa', 'Tarifa'),
            ('valor_retenido', 'Valor Retenido'),
            ('cuenta_pasivo', 'Cuenta Pasivo'),
            ('periodo', 'Período')
        ]
        
        # Filtrar columnas existentes
        columnas_existentes = [(orig, dest) for orig, dest in columnas if orig in df.columns]
        df_detalle = df[[orig for orig, _ in columnas_existentes]].copy()
        df_detalle.columns = [dest for _, dest in columnas_existentes]
        
        # Escribir encabezados
        for col_idx, col_name in enumerate(df_detalle.columns, 1):
            cell = ws_detalle.cell(row=1, column=col_idx, value=col_name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal='center')
        
        # Escribir datos
        for r_idx, row in enumerate(df_detalle.itertuples(index=False), 2):
            for c_idx, value in enumerate(row, 1):
                cell = ws_detalle.cell(row=r_idx, column=c_idx, value=value)
                if c_idx in [7, 8, 9]:  # Columnas numéricas
                    cell.number_format = '#,##0.00'
        
        # Ajustar ancho de columnas
        for col in ws_detalle.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws_detalle.column_dimensions[column].width = adjusted_width
        
        # --- Hoja 2: Resumen por Concepto ---
        ws_resumen = wb.create_sheet("Resumen por Concepto")
        
        # Agrupar por concepto
        df_resumen = df.groupby('concepto_dian').agg({
            'valor_retenido': ['count', 'sum'],
            'base_gravable': 'sum'
        }).round(2)
        
        # Renombrar columnas
        df_resumen.columns = ['Cantidad', 'Total Retenido', 'Total Base']
        df_resumen = df_resumen.reset_index()
        df_resumen.columns = ['Concepto DIAN', 'Cantidad', 'Total Retenido', 'Total Base']
        
        # Agregar nombre del concepto
        from app.core.constants import CONCEPTOS_RETENCION_DIAN
        df_resumen['Concepto'] = df_resumen['Concepto DIAN'].map(
            lambda x: CONCEPTOS_RETENCION_DIAN.get(str(x), x)
        )
        
        # Reordenar columnas
        df_resumen = df_resumen[['Concepto DIAN', 'Concepto', 'Cantidad', 'Total Base', 'Total Retenido']]
        
        # Escribir resumen
        for col_idx, col_name in enumerate(df_resumen.columns, 1):
            cell = ws_resumen.cell(row=1, column=col_idx, value=col_name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
            cell.alignment = Alignment(horizontal='center')
        
        for r_idx, row in enumerate(df_resumen.itertuples(index=False), 2):
            for c_idx, value in enumerate(row, 1):
                cell = ws_resumen.cell(row=r_idx, column=c_idx, value=value)
                if c_idx in [4, 5]:  # Columnas numéricas
                    cell.number_format = '#,##0.00'
        
        # Ajustar ancho
        for col in ws_resumen.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws_resumen.column_dimensions[column].width = adjusted_width
        
        # --- Hoja 3: Resumen por Período ---
        ws_periodo = wb.create_sheet("Resumen por Período")
        
        if 'periodo' in df.columns:
            df_periodo = df.groupby('periodo').agg({
                'valor_retenido': ['count', 'sum']
            }).round(2)
            df_periodo.columns = ['Cantidad', 'Total Retenido']
            df_periodo = df_periodo.reset_index()
            df_periodo.columns = ['Período', 'Cantidad', 'Total Retenido']
            
            for col_idx, col_name in enumerate(df_periodo.columns, 1):
                cell = ws_periodo.cell(row=1, column=col_idx, value=col_name)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
                cell.alignment = Alignment(horizontal='center')
            
            for r_idx, row in enumerate(df_periodo.itertuples(index=False), 2):
                for c_idx, value in enumerate(row, 1):
                    cell = ws_periodo.cell(row=r_idx, column=c_idx, value=value)
                    if c_idx == 3:
                        cell.number_format = '#,##0.00'
            
            for col in ws_periodo.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 30)
                ws_periodo.column_dimensions[column].width = adjusted_width
        
        # Guardar archivo
        wb.save(file_path)
        
        return file_path
    
    def generar_reporte_exogena(
        self,
        datos: List[Dict],
        formato: str = "1001",
        periodo: Optional[str] = None
    ) -> Path:
        """
        Generar reporte de exógena en Excel
        
        Args:
            datos: Lista de datos procesados
            formato: Formato DIAN (1001, 1007, 1008, 1009)
            periodo: Período (opcional)
        
        Returns:
            Path del archivo generado
        """
        if not datos:
            raise ValueError("No hay datos para generar")
        
        fecha = datetime.now().strftime("%Y%m%d")
        periodo_str = f"_{periodo}" if periodo else ""
        filename = f"{fecha}_Exogena_{formato}{periodo_str}.xlsx"
        file_path = self.output_dir / filename
        
        # Crear DataFrame
        df = pd.DataFrame(datos)
        
        # Crear workbook
        wb = Workbook()
        ws = wb.active
        ws.title = f"Formato {formato}"
        
        # Renombrar columnas
        column_mapping = {
            'nit_tercero': 'NIT',
            'nombre_tercero': 'Nombre Tercero',
            'valor': 'Valor',
            'codigo_puc': 'Código PUC',
            'concepto_asignado': 'Concepto DIAN',
            'formato_asignado': 'Formato',
            'estado': 'Estado'
        }
        
        df = df.rename(columns=column_mapping)
        
        # Filtrar columnas existentes
        columnas = [col for col in column_mapping.values() if col in df.columns]
        df = df[columnas]
        
        # Escribir encabezados
        for col_idx, col_name in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal='center')
        
        # Escribir datos
        for r_idx, row in enumerate(df.itertuples(index=False), 2):
            for c_idx, value in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=value)
        
        # Ajustar ancho
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width
        
        wb.save(file_path)
        return file_path